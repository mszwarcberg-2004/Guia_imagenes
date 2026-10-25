"""Validacion externa del TP PSIB: procesa un video nuevo, sin seleccion manual.

COMO USAR (para la catedra):
  1. Cambiar la ruta del video en la variable VIDEO_PATH (esta marcada mas abajo
     con un cartel). Es lo unico que hay que cambiar.
  2. Desde la carpeta raiz del repositorio, ejecutar:
         python3 src/probar_video_externo.py

  Alternativa, sin editar nada: pasar la ruta al ejecutar.
         python3 src/probar_video_externo.py ruta/al/video.mov

SALIDA, en resultados_externos/<nombre_del_video>/ :
  resultados.xlsx     UN solo archivo con dos hojas:
                        - Parametros: latencia, amplitud, velocidad de constriccion y de
                          redilatacion (cada uno una sola vez), mas datos del registro
                        - Curva: area y diametro de la pupila frame a frame
  curva_pupilar.png   la misma curva, como grafico
  extras/             material de apoyo (video de revision, log, tablas internas)

El seguimiento usa el metodo mas robusto de nuestras pruebas (Snakes, sin morfologia).
Si se vuelve a procesar un video, su carpeta de resultados se reemplaza por completo.
"""

from pathlib import Path
import contextlib
import io
import shutil
import sys
import traceback
import warnings

import cv2
import numpy as np
import pandas as pd
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt

import funciones_modulo3
from funciones_modulo3 import procesar_registros


RAIZ = Path(__file__).resolve().parent.parent

# ==============================================================================
#   >>>   LO UNICO QUE HAY QUE CAMBIAR: LA RUTA DEL VIDEO   <<<
#
#   Poner entre comillas la ruta del video. Puede ser absoluta o relativa a la
#   carpeta del repositorio. Ejemplos:
#       VIDEO_PATH = "/Users/nombre/Descargas/video_catedra.mov"
#       VIDEO_PATH = "datos/Videos_ojo/otro_video.mov"
# ==============================================================================
VIDEO_PATH = "datos/Videos_ojo/video_prueba.mp4"
# ==============================================================================

# Metodo mas robusto de nuestras pruebas (Modulo 2: mayor Dice contra la referencia manual)
METODO = "Snakes"
MORFOLOGIA = "Sin morfologia"

SALIDAS = RAIZ / "resultados_externos"

EXTENSIONES_VIDEO = {".mov", ".mp4", ".avi", ".mkv", ".m4v"}

# Columnas de la hoja "Curva" (una fila por frame)
COLUMNAS_CURVA = [
    "frame",
    "tiempo_s",
    "area_px2",
    "diametro_px",
    "area_original_px2",
    "diametro_original_px",
    "detectada"
]

NOMBRES_DE_RESULTADO = ("resultados.xlsx", "parametros.csv")


# ============================================================
# GRAFICO Y VIDEO DE REVISION
# ============================================================

def guardar_curva(serie, destino):

    fig, axes = plt.subplots(
        2, 1,
        figsize=(11, 7),
        sharex=True
    )

    t = serie["tiempo_s"]

    axes[0].plot(
        t,
        serie["area_original_px2"],
        alpha=0.3,
        label="Area detectada"
    )

    axes[0].plot(
        t,
        serie["area_procesada_px2"],
        label="Area procesada"
    )

    axes[0].set_ylabel("Area (px²)")

    axes[1].plot(
        t,
        serie["diametro_original_px"],
        alpha=0.3,
        label="Diametro detectado"
    )

    axes[1].plot(
        t,
        serie["diametro_procesado_px"],
        label="Diametro procesado"
    )

    axes[1].set_ylabel(
        "Diametro equivalente (px)"
    )

    axes[1].set_xlabel("Tiempo (s)")

    for ax in axes:
        ax.grid(alpha=0.25)
        ax.legend()

    fig.suptitle(
        "Seguimiento pupilar automatico"
    )

    fig.tight_layout()
    fig.savefig(destino, dpi=140)
    plt.close(fig)


def guardar_video_revision(video_path, mediciones, destino):

    """
    Dibuja el centro y el circulo equivalente.

    IMPORTANTE: No es el contorno segmentado exacto.
    """

    cap = cv2.VideoCapture(str(video_path))

    if not cap.isOpened():
        raise RuntimeError(
            "No se pudo abrir el video para la vista previa"
        )

    fps = cap.get(cv2.CAP_PROP_FPS)

    ancho = int(
        cap.get(cv2.CAP_PROP_FRAME_WIDTH)
    )

    alto = int(
        cap.get(cv2.CAP_PROP_FRAME_HEIGHT)
    )

    if (
        not np.isfinite(fps)
        or fps <= 0
        or ancho <= 0
        or alto <= 0
    ):
        cap.release()
        raise RuntimeError(
            "FPS o dimensiones invalidas"
        )

    writer = cv2.VideoWriter(
        str(destino),
        cv2.VideoWriter_fourcc(*"mp4v"),
        fps,
        (ancho, alto)
    )

    if not writer.isOpened():
        cap.release()
        raise RuntimeError(
            "No fue posible crear el video MP4"
        )

    datos = mediciones.set_index(
        "frame",
        drop=False
    )

    idx = 0

    try:

        while True:

            ok, frame = cap.read()

            if not ok:
                break

            if idx in datos.index:

                m = datos.loc[idx]

                if isinstance(m, pd.DataFrame):
                    m = m.iloc[0]

                valido = (
                    int(m["detectada"]) == 1
                    and all(
                        pd.notna(m[k])
                        for k in (
                            "centro_x",
                            "centro_y",
                            "diametro_px"
                        )
                    )
                )

                if valido:

                    centro = (
                        int(round(m["centro_x"])),
                        int(round(m["centro_y"]))
                    )

                    radio = max(
                        1,
                        int(
                            round(
                                float(m["diametro_px"]) / 2
                            )
                        )
                    )

                    cv2.circle(
                        frame,
                        centro,
                        radio,
                        (0, 255, 0),
                        2
                    )

                    cv2.circle(
                        frame,
                        centro,
                        2,
                        (0, 255, 0),
                        -1
                    )

                    rotulo = (
                        f"Frame {idx}: "
                        f"area {m['area_px2']:.0f} px2"
                    )

                else:

                    rotulo = (
                        f"Frame {idx}: SIN DETECCION"
                    )

            else:

                rotulo = (
                    f"Frame {idx}: SIN MEDICION"
                )

            cv2.rectangle(
                frame,
                (0, 0),
                (min(ancho, 530), 33),
                (0, 0, 0),
                -1
            )

            cv2.putText(
                frame,
                rotulo,
                (8, 23),
                cv2.FONT_HERSHEY_SIMPLEX,
                0.62,
                (255, 255, 255),
                2
            )

            writer.write(frame)

            idx += 1

    finally:
        cap.release()
        writer.release()

# ============================================================
# ARCHIVO DE RESULTADOS: PARAMETROS (UNA VEZ) + CURVA
# ============================================================

def construir_curva(serie, mediciones):
    """Curva frame a frame: area y diametro de la pupila."""

    curva = pd.DataFrame({
        "frame": mediciones["frame"].to_numpy(),
        "tiempo_s": serie["tiempo_s"].to_numpy(),
        "area_px2": serie["area_procesada_px2"].to_numpy(),
        "diametro_px": serie["diametro_procesado_px"].to_numpy(),
        "area_original_px2": serie["area_original_px2"].to_numpy(),
        "diametro_original_px": serie["diametro_original_px"].to_numpy(),
        "detectada": mediciones["detectada"].to_numpy()
    })

    return curva[COLUMNAS_CURVA]


def construir_parametros(parametros, curva, tipo_registro, video):
    """Parametros fisiologicos, una sola vez.

    Devuelve (principales, registro): listas de (nombre, valor, unidad).
    Un valor None significa que no se pudo calcular.
    """

    p = parametros.iloc[0]

    def numero(campo):
        v = pd.to_numeric(p.get(campo, np.nan), errors="coerce")
        return float(v) if pd.notna(v) else None

    basal = numero("area_basal_px2")
    amplitud = numero("amplitud_px2")

    amplitud_pct = (
        100.0 * amplitud / basal
        if basal is not None and amplitud is not None and basal > 0
        else None
    )

    principales = [
        ("Latencia", numero("latencia_s"), "s"),
        ("Amplitud de constricción", amplitud, "px²"),
        ("Amplitud de constricción (% del área basal)", amplitud_pct, "%"),
        ("Velocidad de constricción", numero("velocidad_constriccion_px2_s"), "px²/s"),
        ("Velocidad de redilatación", numero("velocidad_redilatacion_px2_s"), "px²/s"),
        ("Estado del cálculo", str(p.get("estado", "")), "")
    ]

    registro = [
        ("Video", video.name, ""),
        ("Método de segmentación", f"{METODO} + {MORFOLOGIA}", ""),
        ("Tipo de registro (A: la luz se enciende, B: la luz se apaga)", tipo_registro or None, ""),
        ("Encendido de la luz", numero("tiempo_estimulo_s"), "s"),
        ("Apagado de la luz", numero("tiempo_apagado_s"), "s"),
        ("Área basal (antes del estímulo)", basal, "px²"),
        ("Área mínima (durante la constricción)", numero("area_minima_px2"), "px²"),
        ("Frames totales", int(len(curva)), ""),
        ("Frames con pupila detectada", int(curva["detectada"].sum()), "")
    ]

    return principales, registro


def _texto_valor(valor):
    """Valor para mostrar en la hoja de parametros."""

    if valor is None:
        return "no calculado"

    if isinstance(valor, float):
        return round(valor, 3)

    return valor


def guardar_resultados(carpeta, principales, registro, curva):
    """Guarda resultados.xlsx (hojas Parametros y Curva).

    Si no esta instalado openpyxl, guarda parametros.csv y curva.csv.
    Devuelve los nombres de los archivos creados.
    """

    curva_redondeada = curva.round(4)

    try:

        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

    except ImportError:

        filas = [
            (nombre, _texto_valor(valor), unidad)
            for nombre, valor, unidad in principales + registro
        ]

        pd.DataFrame(
            filas, columns=["Parametro", "Valor", "Unidad"]
        ).to_csv(carpeta / "parametros.csv", index=False, encoding="utf-8-sig")

        curva_redondeada.to_csv(carpeta / "curva.csv", index=False)

        print(
            "\nAviso: no esta instalado openpyxl, por eso se guardaron "
            "parametros.csv y curva.csv.\n"
            "Para obtener un unico archivo Excel: pip install openpyxl"
        )

        return ["parametros.csv", "curva.csv"]

    negrita = Font(bold=True)
    relleno = PatternFill("solid", fgColor="DDEBF7")

    libro = Workbook()

    # ---- Hoja 1: parametros ----
    hoja = libro.active
    hoja.title = "Parametros"

    hoja.append(["Parámetro", "Valor", "Unidad"])

    for celda in hoja[1]:
        celda.font = negrita
        celda.fill = relleno

    for nombre, valor, unidad in principales:
        hoja.append([nombre, _texto_valor(valor), unidad])
        hoja.cell(row=hoja.max_row, column=1).font = negrita

    hoja.append([])
    hoja.append(["Datos del registro", "", ""])

    for celda in hoja[hoja.max_row]:
        celda.font = negrita
        celda.fill = relleno

    for nombre, valor, unidad in registro:
        hoja.append([nombre, _texto_valor(valor), unidad])

    hoja.column_dimensions["A"].width = 58
    hoja.column_dimensions["B"].width = 44
    hoja.column_dimensions["C"].width = 8

    # ---- Hoja 2: curva ----
    hoja_curva = libro.create_sheet("Curva")

    hoja_curva.append(list(curva_redondeada.columns))

    for celda in hoja_curva[1]:
        celda.font = negrita
        celda.fill = relleno

    filas = (
        curva_redondeada
        .astype(object)
        .where(curva_redondeada.notna(), None)
        .values
        .tolist()
    )

    for fila in filas:
        hoja_curva.append(fila)

    hoja_curva.freeze_panes = "A2"

    for letra in "ABCDEFG":
        hoja_curva.column_dimensions[letra].width = 20

    libro.save(carpeta / "resultados.xlsx")

    return ["resultados.xlsx"]


# ============================================================
# EJECUCION PRINCIPAL
# ============================================================

def ejecutar(video_path):

    video = Path(video_path).expanduser().resolve()

    if not video.is_file():
        raise FileNotFoundError(
            f"No existe el video: {video}"
        )

    # Carpeta de salida: se reemplaza por completo para no mezclar corridas
    carpeta = SALIDAS / video.stem

    if carpeta.exists():
        shutil.rmtree(carpeta)

    extras = carpeta / "extras"
    resultados = extras / "resultados"
    resultados.mkdir(parents=True)

    print(f"Video seleccionado: {video}")
    print(f"Metodo de segmentacion: {METODO} + {MORFOLOGIA}")

    # ========================================================
    # 1. SEGUIMIENTO DE LA PUPILA Y PARAMETROS
    # ========================================================

    print("\nSiguiendo la pupila y calculando parametros...")

    # Este script lee los archivos con los nombres del formato anterior
    funciones_modulo3.GUARDAR_FORMATO_ANTERIOR = True

    log = io.StringIO()

    try:

        with contextlib.redirect_stdout(log):
            with warnings.catch_warnings():

                warnings.simplefilter("ignore", RuntimeWarning)

                procesar_registros(
                    video,
                    resultados,
                    METODO,
                    MORFOLOGIA,
                    modo="evaluacion",
                    procesar_otros=False,
                    mostrar_figuras=False
                )

    finally:

        plt.close("all")

        (extras / "ejecucion.txt").write_text(
            log.getvalue(),
            encoding="utf-8"
        )

    # ========================================================
    # 2. ARCHIVO DE RESULTADOS
    # ========================================================

    serie_csv = resultados / f"serie_procesada_{video.stem}.csv"
    mediciones_csv = resultados / f"mediciones_{video.stem}.csv"
    parametros_csv = resultados / f"parametros_fisiologicos_{video.stem}.csv"

    for archivo in [serie_csv, mediciones_csv, parametros_csv]:
        if not archivo.exists():
            raise RuntimeError(
                f"No se genero {archivo.name}. Ver extras/ejecucion.txt"
            )

    serie = pd.read_csv(serie_csv)
    mediciones = pd.read_csv(mediciones_csv)
    parametros = pd.read_csv(parametros_csv)

    if serie.empty or mediciones.empty or parametros.empty:
        raise RuntimeError("Se genero un CSV vacio")

    if len(serie) != len(mediciones):
        raise RuntimeError(
            "La serie y las mediciones tienen distinta cantidad de frames"
        )

    # Tipo de registro: A (la luz se enciende) o B (la luz se apaga)
    tipo_registro = ""

    try:

        limpio = pd.read_csv(
            resultados / f"parametros_{video.stem}.csv",
            encoding="utf-8-sig"
        )

        tipo_registro = str(limpio.iloc[0]["tipo_registro"])

    except Exception:
        pass

    curva = construir_curva(serie, mediciones)

    principales, registro = construir_parametros(
        parametros, curva, tipo_registro, video
    )

    archivos = guardar_resultados(carpeta, principales, registro, curva)

    guardar_curva(
        serie,
        carpeta / "curva_pupilar.png"
    )

    # ========================================================
    # 3. VIDEO DE REVISION (si falla, solo se avisa)
    # ========================================================

    print("\nGenerando video de revision...")

    try:

        guardar_video_revision(
            video,
            mediciones,
            extras / "revision_pupila.mp4"
        )

    except Exception as exc:

        print(f"No se pudo generar el video de revision: {exc}")

    # ========================================================
    # 4. RESUMEN EN PANTALLA
    # ========================================================

    print("\n===== RESULTADOS =====")

    for nombre, valor, unidad in principales:

        texto = _texto_valor(valor)

        print(f"{nombre}: {texto} {unidad}".rstrip())

    print(f"Tipo de registro: {tipo_registro or 'no determinado'}")

    print(
        f"Frames con pupila detectada: "
        f"{int(curva['detectada'].sum())}/{len(curva)}"
    )

    print(f"\nResultados guardados en: {carpeta}")

    print(
        f"  {' y '.join(archivos)}   parametros y curva (archivo principal)\n"
        "  curva_pupilar.png   la curva como grafico\n"
        "  extras/             video de revision, log y tablas internas"
    )

    print(
        "IMPORTANTE: revisar visualmente el seguimiento "
        "(extras/revision_pupila.mp4) antes de interpretar "
        "los parametros fisiologicos."
    )


# ============================================================
# BUSQUEDA DE VIDEOS E INICIO DEL PROGRAMA
# ============================================================

def buscar_videos():
    """Videos a procesar: los de la terminal (archivos o carpetas) o VIDEO_PATH."""

    argumentos = sys.argv[1:]

    if argumentos:

        rutas = [Path(a).expanduser() for a in argumentos]

    else:

        ruta = Path(VIDEO_PATH).expanduser()

        # Una ruta relativa se interpreta desde la carpeta del repositorio
        if not ruta.is_absolute():
            ruta = RAIZ / ruta

        rutas = [ruta]

    videos = []

    for ruta in rutas:

        if ruta.is_dir():

            videos += sorted(
                p for p in ruta.iterdir()
                if p.is_file() and p.suffix.lower() in EXTENSIONES_VIDEO
            )

        else:

            # Si no existe, ejecutar() avisa con un error claro
            videos.append(ruta)

    return videos


if __name__ == "__main__":

    videos = buscar_videos()

    if not videos:

        print("No se encontraron videos para procesar en la carpeta indicada.")

        raise SystemExit(1)

    errores = []

    for video in videos:

        print(f"\n{'=' * 60}\nVIDEO: {video.name}\n{'=' * 60}")

        try:

            ejecutar(video)

        except Exception as exc:

            errores.append((video.name, str(exc)))

            print(f"ERROR EN {video.name}: {exc}")

            traceback.print_exc()

            # No dejar una carpeta incompleta de un video que fallo
            fallida = SALIDAS / video.stem

            if (
                fallida.exists()
                and not any((fallida / n).exists() for n in NOMBRES_DE_RESULTADO)
            ):
                shutil.rmtree(fallida, ignore_errors=True)

    if errores:

        print("\nNo se pudieron procesar:")

        for nombre, motivo in errores:
            print(f"  - {nombre}: {motivo}")

        if any(m.startswith("No existe el video") for _, m in errores):
            print(
                "\nCambiar la ruta de VIDEO_PATH en src/probar_video_externo.py, "
                "o pasarla al ejecutar:\n"
                "  python3 src/probar_video_externo.py ruta/al/video.mov"
            )

        raise SystemExit(1)